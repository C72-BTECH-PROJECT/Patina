"""Generate the synthetic resume fixtures used by the T-0001 harness.

Every person, employer, school, email, phone number and link below is
fictional (example.com / example-* handles). Do not put real resumes here;
those go in tests/private/ (git-ignored).

Run with the NLP venv (needs python-docx, fpdf2, Pillow):
    nlp-engine/venv/Scripts/python tests/fixtures/make_fixtures.py

Output is byte-stable across runs (fixed PDF/DOCX timestamps) so regenerating
does not show up as a diff.
"""

import datetime as dt
import io
import zipfile
from pathlib import Path

from docx import Document
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
FIXED_TS = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc)


# --------------------------------------------------------------------------
# Resume content. Each resume is a list of (heading, lines) sections plus a
# contact header, so the same data can be written to PDF, DOCX or an image.
# --------------------------------------------------------------------------

F1 = {
    "name": "Maya Okafor",
    "contact": [
        "maya.okafor@example.com | +1 555 010 4477 | Austin, TX",
        "github.com/example-mokafor | linkedin.com/in/example-maya-okafor",
    ],
    "sections": [
        ("Summary", [
            "Backend-leaning software engineer with 3 years of experience building",
            "Python web services and React front ends.",
        ]),
        ("Skills", [
            "Python, Django, FastAPI, React, JavaScript, PostgreSQL, SQL, Docker,",
            "AWS, Git, REST API, Redis, Pytest",
        ]),
        ("Experience", [
            "Software Engineer - Northwind Labs (Jun 2022 - Present)",
            "- Built REST APIs in Django and FastAPI serving 40k daily users.",
            "- Cut p95 latency by 35% by adding Redis caching and query indexes.",
            "- Containerised services with Docker and deployed to AWS ECS.",
            "Software Engineering Intern - Bluebird Analytics (Jan 2022 - May 2022)",
            "- Wrote ETL jobs in Python and SQL for the reporting warehouse.",
        ]),
        ("Projects", [
            "Ledgerly - personal finance tracker (React, FastAPI, PostgreSQL)",
            "github.com/example-mokafor/ledgerly",
            "QueueBench - load-testing harness for job queues (Python, Redis)",
        ]),
        ("Education", [
            "B.Tech in Computer Science - Example Institute of Technology (2018 - 2022)",
        ]),
    ],
}

F2 = {
    "name": "Arjun Mehta",
    "contact": [
        "arjun.mehta@example.com | +91 98765 01234 | Pune, India",
        "github.com/example-arjunm",
    ],
    "sections": [
        ("Summary", [
            "Java and Node.js developer with 4 years in product companies.",
        ]),
        ("Technical Skills", [
            "Java, Spring Boot, Node.js, Express, MongoDB, MySQL, Kubernetes,",
            "Docker, Git, Jenkins, Microservices",
        ]),
        ("Work Experience", [
            "Senior Developer - Contoso Retail Tech (Aug 2021 - Present)",
            "- Split a monolith into six Spring Boot microservices on Kubernetes.",
            "- Owned the order service: 99.95% availability over 12 months.",
            "Developer - Fabrikam Software (Jul 2019 - Jul 2021)",
            "- Built Express APIs backed by MongoDB for a logistics dashboard.",
        ]),
        ("Projects", [
            "RateLimiter - token-bucket middleware for Express",
            "github.com/example-arjunm/ratelimiter",
        ]),
        ("Education", [
            "B.E. in Information Technology - Example University of Pune (2015 - 2019)",
        ]),
    ],
}

F3 = {
    "name": "Lena Fischer",
    "contact": [
        "lena.fischer@example.com | +49 151 5550 1234 | Hamburg, Germany",
    ],
    "sections": [
        ("Summary", [
            "IT support engineer focused on incident management and keeping",
            "production systems healthy. No public code portfolio.",
        ]),
        ("Skills", [
            "Linux, Windows Server, SQL, Bash, ServiceNow, Jira, ITIL,",
            "Nagios, Active Directory, Incident Management, Troubleshooting",
        ]),
        ("Experience", [
            "Application Support Engineer - Tailspin Logistics (Mar 2021 - Present)",
            "- Resolved 600+ L2 tickets with a 96% SLA hit rate.",
            "- Wrote Bash and SQL runbooks that cut mean time to resolve by 30%.",
            "- Ran monthly patching for 120 Linux and Windows servers.",
            "IT Helpdesk Analyst - Woodgrove Bank (Jan 2019 - Feb 2021)",
            "- First-line support for 800 staff; managed Active Directory accounts.",
        ]),
        ("Certifications", [
            "ITIL 4 Foundation (2021)",
        ]),
        ("Education", [
            "B.Sc. in Information Systems - Example University Hamburg (2015 - 2018)",
        ]),
    ],
}

F4 = {
    "name": "Samuel Adeyemi",
    "contact": [
        "samuel.adeyemi@example.com | +44 7700 900123 | Leeds, UK",
    ],
    "sections": [
        ("Profile", [
            "Production support and maintenance engineer. Keeps legacy",
            "billing systems running; most work is internal and not on GitHub.",
        ]),
        ("Core Skills", [
            "SQL, Oracle, PL/SQL, Shell Scripting, Python, Splunk,",
            "AWS CloudWatch, Jenkins, Root Cause Analysis, On-call",
        ]),
        ("Employment", [
            "Production Support Engineer - Litware Billing (May 2020 - Present)",
            "- Primary on-call for a 15-year-old Oracle billing platform.",
            "- Automated nightly health checks in Python and Splunk alerts.",
            "- Led 20+ root-cause analyses; reduced repeat incidents by 40%.",
            "Maintenance Engineer - Adventure Works IT (Jun 2017 - Apr 2020)",
            "- Patched and maintained PL/SQL batch jobs and shell scripts.",
        ]),
        ("Education", [
            "HND in Computing - Example College Leeds (2015 - 2017)",
        ]),
    ],
}

F5 = {
    "name": "Rosa Delgado",
    "sidebar": [
        ("Contact", [
            "rosa.delgado@example.com",
            "+1 555 010 9988",
            "Denver, CO",
            "github.com/example-rdelgado",
        ]),
        ("Skills", [
            "Python", "FastAPI", "PostgreSQL", "SQL", "Docker", "Kubernetes",
            "AWS", "Terraform", "Git",
        ]),
        ("Education", [
            "B.S. Computer Science",
            "Example State University",
            "2016 - 2020",
        ]),
    ],
    "main": [
        ("Experience", [
            "Platform Engineer",
            "Proseware Cloud (Feb 2022 - Present)",
            "- Ran 40 Kubernetes services on AWS EKS with Terraform.",
            "- Built FastAPI internal tools for deploy approvals.",
            "Backend Developer",
            "Lucerne Media (Jul 2020 - Jan 2022)",
            "- Wrote Python services over PostgreSQL for ad reporting.",
        ]),
        ("Projects", [
            "infra-lint - Terraform policy checker",
            "github.com/example-rdelgado/infra-lint",
        ]),
    ],
}

F6 = {
    "name": "Tomas Novak",
    "contact": [
        "tomas.novak@example.com | +420 555 010 222 | Brno, Czechia",
        "github.com/example-tnovak",
    ],
    "sections": [
        ("Skills", [
            "Python, Django, PostgreSQL, Docker, Git, REST API",
        ]),
        ("Experience", [
            "Python Developer - Margie's Travel (Sep 2021 - Present)",
            "- Maintained Django booking APIs and PostgreSQL schemas.",
        ]),
        ("Education", [
            "M.Sc. Informatics - Example Technical University Brno (2019 - 2021)",
        ]),
    ],
}


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------

def _new_pdf():
    pdf = FPDF(format="A4")
    pdf.set_creation_date(FIXED_TS)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    return pdf


def write_single_column_pdf(data, path):
    pdf = _new_pdf()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, data["name"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for line in data["contact"]:
        pdf.cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    for heading, lines in data["sections"]:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 7, heading.upper(), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        for line in lines:
            pdf.multi_cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def write_two_column_pdf(data, path):
    """Left sidebar (contact/skills/education) beside a main column."""
    pdf = _new_pdf()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, data["name"], new_x="LMARGIN", new_y="NEXT")
    top = pdf.get_y() + 2
    left_x, left_w = 10, 55
    right_x, right_w = 72, 128

    def column(x, w, sections):
        pdf.set_xy(x, top)
        for heading, lines in sections:
            pdf.set_x(x)
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(w, 7, heading.upper(), new_x="LEFT", new_y="NEXT")
            pdf.set_font("Helvetica", "", 10)
            for line in lines:
                pdf.set_x(x)
                pdf.multi_cell(w, 5, line, new_x="LEFT", new_y="NEXT")
            pdf.ln(3)

    column(left_x, left_w, data["sidebar"])
    column(right_x, right_w, data["main"])
    pdf.output(str(path))


def _load_font(size):
    for name in ("arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def write_image_only_pdf(data, path):
    """Render the resume as a bitmap and embed only that image (no text layer),
    the way a scanned resume arrives. Exercises the missing-OCR path."""
    width, height = 1240, 1754  # A4 at 150 dpi
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title, body, head = _load_font(44), _load_font(24), _load_font(28)
    y = 80
    draw.text((80, y), data["name"], fill="black", font=title)
    y += 70
    for line in data["contact"]:
        draw.text((80, y), line, fill="black", font=body)
        y += 34
    for heading, lines in data["sections"]:
        y += 24
        draw.text((80, y), heading.upper(), fill="black", font=head)
        y += 42
        for line in lines:
            draw.text((80, y), line, fill="black", font=body)
            y += 34
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    buf.seek(0)
    pdf = FPDF(format="A4")
    pdf.set_creation_date(FIXED_TS)
    pdf.add_page()
    pdf.image(buf, x=0, y=0, w=210, h=297)
    pdf.output(str(path))


def write_docx(data, path):
    doc = Document()
    props = doc.core_properties
    props.author = "T-0001 fixture generator"
    props.created = FIXED_TS.replace(tzinfo=None)
    props.modified = FIXED_TS.replace(tzinfo=None)
    props.last_modified_by = "T-0001 fixture generator"
    props.revision = 1
    doc.add_heading(data["name"], level=0)
    for line in data["contact"]:
        doc.add_paragraph(line)
    for heading, lines in data["sections"]:
        doc.add_heading(heading, level=1)
        for line in lines:
            if line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    # python-docx stamps each zip entry with the current time; rewrite the
    # archive with a fixed timestamp so the file is byte-stable.
    buf.seek(0)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(item.filename))


def main():
    write_single_column_pdf(F1, HERE / "f1_dev_github.pdf")
    write_docx(F2, HERE / "f2_dev_github.docx")
    write_single_column_pdf(F3, HERE / "f3_support_nogithub.pdf")
    write_docx(F4, HERE / "f4_support_nogithub.docx")
    write_two_column_pdf(F5, HERE / "f5_multicolumn.pdf")
    write_image_only_pdf(F6, HERE / "f6_image_only.pdf")
    print("wrote 6 fixtures to", HERE)


if __name__ == "__main__":
    main()
